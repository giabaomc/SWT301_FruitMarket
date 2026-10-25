import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

sys.stdout.reconfigure(encoding='utf-8')

file_path = r'C:\Users\ganda\Downloads\UnitTest_Login.xlsx'
wb = openpyxl.load_workbook(file_path)

print("Original sheets:", wb.sheetnames)

# Define common styles
navy_fill = PatternFill(start_color="000080", end_color="000080", fill_type="solid")
gray_fill = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")

font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Arial", size=10, bold=True)
font_regular = Font(name="Arial", size=10, bold=False)
font_failed = Font(name="Arial", size=10, bold=True, color="FF0000")

thin_side = Side(style='thin', color='000000')
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

align_center = Alignment(horizontal='center', vertical='center')
align_center_top = Alignment(horizontal='center', vertical='top')
align_rotate90 = Alignment(horizontal='center', vertical='center', textRotation=90)
align_left = Alignment(horizontal='left', vertical='center')
align_left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

# ----------------------------------------------------
# 1. UPDATE COVER SHEET
# ----------------------------------------------------
ws_cover = wb['Cover']
ws_cover['F3'] = 'BaoPGHE204203'
ws_cover['F4'] = '19/10/2026'
ws_cover['B5'] = 'FRUIT_MARKET_UnitTest_Login_Search_v1.1'
ws_cover['F5'] = '1.1'

# Add change record row 10
change_row = 10
ws_cover.cell(change_row, 1, '19/10/2026')
ws_cover.cell(change_row, 2, '1.1')
ws_cover.cell(change_row, 3, 'Function list')
ws_cover.cell(change_row, 4, 'A')
ws_cover.cell(change_row, 5, 'Add unit tests of display and search product function (getAllProduct, searchProductByName)')
ws_cover.cell(change_row, 6, 'FunctionList 2')

for c in range(1, 7):
    cell = ws_cover.cell(change_row, c)
    ref_cell = ws_cover.cell(9, c)
    cell.font = copy(ref_cell.font)
    cell.border = copy(ref_cell.border)
    cell.alignment = copy(ref_cell.alignment)

print("Updated Cover sheet.")

# ----------------------------------------------------
# 2. UPDATE FUNCTIONLIST SHEET
# ----------------------------------------------------
ws_fn = wb['FunctionList']
fn_row = 10
ws_fn.cell(fn_row, 1, 2)
ws_fn.cell(fn_row, 2, 'ProductDAO / ProductJSP')
ws_fn.cell(fn_row, 3, 'Display & Search (getAllProduct, searchProductByName)')
ws_fn.cell(fn_row, 4, 'searchProduct')
ws_fn.cell(fn_row, 5, 'Hiển thị danh sách tất cả sản phẩm (submit=null) và tìm kiếm sản phẩm theo tên (submit="Tìm"). Kết quả chỉ gồm các sản phẩm đang hoạt động (status = 1).')
ws_fn.cell(fn_row, 6, 'The database is connected and contains product data.')

for c in range(1, 7):
    cell = ws_fn.cell(fn_row, c)
    ref_cell = ws_fn.cell(9, c)
    cell.font = copy(ref_cell.font)
    cell.border = copy(ref_cell.border)
    cell.alignment = copy(ref_cell.alignment)

print("Updated FunctionList sheet.")

# ----------------------------------------------------
# 3. UPDATE STATISTICS SHEET
# ----------------------------------------------------
ws_stat = wb['Statistics']

# First, unmerge existing merged cells that will shift down
for m in ['A10:B10', 'A12:B12', 'A13:B13', 'A14:B14', 'A15:B15', 'A16:B16']:
    if m in ws_stat.merged_cells:
        ws_stat.unmerge_cells(m)

# Shift rows 10-16 down to rows 11-17
# Let's read old values and formats of rows 10..16
rows_data = []
for r in range(10, 17):
    row_info = []
    for c in range(1, 10):
        cell = ws_stat.cell(r, c)
        row_info.append({
            'val': cell.value,
            'font': copy(cell.font),
            'border': copy(cell.border),
            'fill': copy(cell.fill),
            'alignment': copy(cell.alignment),
            'number_format': cell.number_format
        })
    rows_data.append(row_info)

# Write Row 10: Function 2 (searchProduct)
ws_stat.cell(10, 1, 2)
ws_stat.cell(10, 2, 'searchProduct')
ws_stat.cell(10, 3, '=searchProduct!A5')
ws_stat.cell(10, 4, '=searchProduct!B5')
ws_stat.cell(10, 5, '=searchProduct!C5')
ws_stat.cell(10, 6, '=searchProduct!D5')
ws_stat.cell(10, 7, '=searchProduct!E5')
ws_stat.cell(10, 8, '=searchProduct!F5')
ws_stat.cell(10, 9, '=searchProduct!G5')

for c in range(1, 10):
    cell = ws_stat.cell(10, c)
    ref = ws_stat.cell(9, c)
    cell.font = copy(ref.font)
    cell.border = copy(ref.border)
    cell.alignment = copy(ref.alignment)

# Write Row 11: Sub total
ws_stat.cell(11, 1, 'Sub total')
ws_stat.cell(11, 2, None)
ws_stat.cell(11, 3, '=SUM(C9:C10)')
ws_stat.cell(11, 4, '=SUM(D9:D10)')
ws_stat.cell(11, 5, '=SUM(E9:E10)')
ws_stat.cell(11, 6, '=SUM(F9:F10)')
ws_stat.cell(11, 7, '=SUM(G9:G10)')
ws_stat.cell(11, 8, '=SUM(H9:H10)')
ws_stat.cell(11, 9, '=SUM(I9:I10)')

old_subtotal = rows_data[0] # row 10 previously
for c in range(1, 10):
    cell = ws_stat.cell(11, c)
    cell.font = copy(old_subtotal[c-1]['font'])
    cell.border = copy(old_subtotal[c-1]['border'])
    cell.fill = copy(old_subtotal[c-1]['fill'])
    cell.alignment = copy(old_subtotal[c-1]['alignment'])

# Row 12 is empty blank separator
for c in range(1, 10):
    ws_stat.cell(12, c).value = None

# Row 13 to 17: coverage rows
coverage_formulas = [
    ('Test coverage', '=IF(I11=0,0,(C11+D11)/I11)'),
    ('Test successful coverage', '=IF(I11=0,0,C11/I11)'),
    ('Normal case', '=IF(I11=0,0,F11/I11)'),
    ('Abnormal case', '=IF(I11=0,0,G11/I11)'),
    ('Boundary case', '=IF(I11=0,0,H11/I11)')
]

for idx, (label, formula) in enumerate(coverage_formulas):
    r = 13 + idx
    old_row_info = rows_data[2 + idx] # old row 12..16
    ws_stat.cell(r, 1, label)
    ws_stat.cell(r, 2, None)
    ws_stat.cell(r, 3, formula)
    for c in range(4, 10):
        ws_stat.cell(r, c).value = None
    
    for c in range(1, 10):
        cell = ws_stat.cell(r, c)
        cell.font = copy(old_row_info[c-1]['font'])
        cell.border = copy(old_row_info[c-1]['border'])
        cell.fill = copy(old_row_info[c-1]['fill'])
        cell.alignment = copy(old_row_info[c-1]['alignment'])
        cell.number_format = old_row_info[c-1]['number_format']

# Re-merge cells in Statistics
ws_stat.merge_cells('A11:B11')
ws_stat.merge_cells('A13:B13')
ws_stat.merge_cells('A14:B14')
ws_stat.merge_cells('A15:B15')
ws_stat.merge_cells('A16:B16')
ws_stat.merge_cells('A17:B17')

print("Updated Statistics sheet.")

# ----------------------------------------------------
# 4. CREATE / RECREATE searchProduct SHEET
# ----------------------------------------------------
if 'searchProduct' in wb.sheetnames:
    del wb['searchProduct']

ws = wb.create_sheet(title='searchProduct')

# Set column widths
ws.column_dimensions['A'].width = 11.71
ws.column_dimensions['B'].width = 40.71
ws.column_dimensions['C'].width = 30.71
for col_idx in range(4, 25): # D to X (cols 4 to 24)
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = 5.71
ws.column_dimensions['Y'].width = 9.14

# Set row heights
ws.row_dimensions[3].height = 39.95
ws.row_dimensions[7].height = 62.1
ws.row_dimensions[55].height = 62.1
ws.row_dimensions[56].height = 62.1

# Top Information (Rows 1-5)
ws['A1'] = 'Code Module'
ws['B1'] = 'ProductDAO (called by ProductJSP, service=listProduct)'
ws['C1'] = 'Method'
ws['D1'] = 'Vector<Product> getAllProduct(String sql) / Vector<Product> searchProductByName(String productName)'
ws.merge_cells('D1:X1')

ws['A2'] = 'Created By'
ws['B2'] = 'BaoPGHE204203'
ws['C2'] = 'Executed By'
ws['D2'] = 'BaoPGHE204203'
ws.merge_cells('D2:X2')

ws['A3'] = 'Lines of code'
ws['B3'] = 39
ws['C3'] = 'Test requirement'
ws['D3'] = 'Hiển thị danh sách sản phẩm; tìm kiếm theo tên đầy đủ hoặc một phần; tìm tên không tồn tại; kiểm tra kết quả chỉ gồm sản phẩm đang hoạt động (status = 1). ProductDAO trả về danh sách sản phẩm phù hợp; ProductJSP chuyển tiếp đến productjsp.jsp hiển thị bảng hoa quả hoặc thông báo "Không tìm thấy hoa quả nào.". '
ws.merge_cells('D3:X3')

ws['A4'] = 'Passed'
ws['B4'] = 'Failed'
ws['C4'] = 'Untested'
ws['D4'] = 'N/A/B'
ws.merge_cells('D4:F4')
ws['G4'] = 'Total Test Cases'
ws.merge_cells('G4:X4')

ws['A5'] = '=COUNTIF($D$54:$X$54,"P")'
ws['B5'] = '=COUNTIF($D$54:$X$54,"F")'
ws['C5'] = '=G5-A5-B5'
ws['D5'] = '=COUNTIF($D$53:$X$53,"N")'
ws['E5'] = '=COUNTIF($D$53:$X$53,"A")'
ws['F5'] = '=COUNTIF($D$53:$X$53,"B")'
ws['G5'] = '=COUNTA($D$7:$X$7)'
ws.merge_cells('G5:X5')

# Styling Rows 1-5
for r in [1, 2, 3]:
    ws.cell(r, 1).font = font_bold
    ws.cell(r, 2).font = font_regular
    ws.cell(r, 3).font = font_bold
    ws.cell(r, 4).font = font_regular
    ws.cell(r, 4).alignment = align_left_wrap

for c in range(1, 25):
    cell4 = ws.cell(4, c)
    cell4.fill = gray_fill
    cell4.font = font_bold
    cell4.border = thin_border
    cell4.alignment = align_center

for c in range(1, 25):
    cell5 = ws.cell(5, c)
    cell5.font = font_regular
    cell5.border = thin_border
    cell5.alignment = align_center

# Row 7: UTCID Header
ws['A7'] = None
ws.merge_cells('A7:C7')
for c in range(1, 4):
    cell = ws.cell(7, c)
    cell.fill = navy_fill
    cell.border = thin_border

for i in range(1, 22):
    c = 3 + i # Col 4 is D (UTCID01)
    utcid = f'UTCID{i:02d}'
    cell = ws.cell(7, c, utcid)
    cell.fill = navy_fill
    cell.font = font_header
    cell.alignment = align_rotate90
    cell.border = thin_border

# Row definitions for Condition, Confirm, Result
# Format: (row_num, col_b_val, col_c_val, [list of 'O' column indices 1-based (4=D, etc.)])
condition_rows = [
    (8, 'Precondition', None, []),
    (9, 'Connected to the database (seed data loaded)', None, list(range(4, 24))), # UTCID01-20
    (10, 'Database connection unavailable', None, [24]), # UTCID21
    (11, 'Database contains inactive product (status = 0)', None, [5, 14]), # UTCID02, UTCID11
    (12, 'submit (action / mode)', None, []),
    (13, None, 'null (Display product list)', [4, 5]), # UTCID01, UTCID02
    (14, None, '"Tìm" (Search products)', list(range(6, 25))), # UTCID03-21
    (15, 'productName', None, []),
    (16, 'exact full name', 'Tao Fuji', [6, 16, 24]), # UTCID03, UTCID13, UTCID21
    (17, 'exact full name (Vietnamese)', 'Cam sanh mien Tay', [7]), # UTCID04
    (18, 'case-insensitive', 'tao fuji', [8]), # UTCID05
    (19, 'partial name (suffix/middle)', 'cat Hoa Loc', [9]), # UTCID06
    (20, 'partial name (prefix)', 'Dua', [10]), # UTCID07
    (21, 'partial substring (multiple matches)', 'an', [11]), # UTCID08
    (22, 'not exist in database', 'Sau rieng Cai Mon', [12]), # UTCID09
    (23, 'not exist (random string)', 'xyz999', [13]), # UTCID10
    (24, 'name of deactivated product (status = 0)', 'Dua hau cu (Deactive)', [14]), # UTCID11
    (25, 'empty string', '""', [15]), # UTCID12
    (26, 'with leading/trailing space', '" Tao Fuji "', [16]), # UTCID13
    (27, 'whitespace only', '"   "', [17]), # UTCID14
    (28, 'SQL wildcard %', '"%"', [18]), # UTCID15
    (29, 'SQL injection', "' OR '1'='1", [19]), # UTCID16
    (30, 'XSS script injection', '<script>alert(1)</script>', [20]), # UTCID17
    (31, '150 characters (max length)', 'aaa…a', [21]), # UTCID18
    (32, '151 characters', 'aaa…a', [22]), # UTCID19
    (33, 'null parameter', 'null', [23]) # UTCID20
]

confirm_rows = [
    (34, 'Return (Vector<Product>)', None, []),
    (35, None, 'size = 10 (all active products)', [4, 15, 18]), # UTCID01, UTCID12, UTCID15
    (36, None, 'size = 1 ("Tao Fuji")', [6, 8]), # UTCID03, UTCID05
    (37, None, 'size = 1 ("Cam sanh mien Tay")', [7]), # UTCID04
    (38, None, 'size = 1 ("Xoai cat Hoa Loc")', [9]), # UTCID06
    (39, None, 'size = 1 ("Dua hau khong hat")', [10]), # UTCID07
    (40, None, 'size = 4 (multiple products containing \'an\')', [11]), # UTCID08
    (41, None, 'size = 11 (contains inactive product status = 0)', [5]), # UTCID02 (Defect SEARCH-D01)
    (42, None, 'size = 0 (empty vector)', [12, 13, 14, 16, 17, 19, 20, 21, 22, 23]), # UTCID09,10,11,13(Defect),14,16,17,18,19,20
    (43, None, 'null (throws exception)', [24]), # UTCID21 (Defect SEARCH-D03)
    (44, 'Exception', None, []),
    (45, None, 'None', list(range(4, 24))), # UTCID01-20
    (46, None, 'NullPointerException', [24]), # UTCID21
    (47, 'Screen message (productjsp.jsp)', None, []),
    (48, 'table displays active products', '10 products', [4, 15, 18]), # UTCID01, UTCID12, UTCID15
    (49, 'table displays matching product(s)', '1 to 4 products', [6, 7, 8, 9, 10, 11]), # UTCID03-08
    (50, 'table displays inactive product', 'product with "DeActive" badge', [5]), # UTCID02 (Defect)
    (51, 'table shows empty state', '"Không tìm thấy hoa quả nào."', [12, 13, 14, 16, 17, 19, 20, 21, 22, 23]), # UTCID09-11, 13, 14, 16-20
    (52, 'server error', 'HTTP Status 500 – Internal Server Error', [24]) # UTCID21
]

# Write Condition rows
ws.merge_cells('A8:A33')
ws['A8'] = 'Condition'
ws['A8'].fill = navy_fill
ws['A8'].font = font_header
ws['A8'].alignment = align_center_top

for r, b_val, c_val, o_cols in condition_rows:
    if b_val:
        ws.cell(r, 2, b_val)
        if r in [8, 12, 15]:
            ws.cell(r, 2).font = font_bold
        else:
            ws.cell(r, 2).font = font_regular
    if c_val:
        ws.cell(r, 3, c_val)
        ws.cell(r, 3).font = font_regular
    
    # Border for cols A..C
    for c in range(1, 4):
        ws.cell(r, c).border = thin_border
    
    # Fill O markers and borders for cols D..X
    for col_idx in range(4, 25):
        cell = ws.cell(r, col_idx)
        cell.border = thin_border
        cell.alignment = align_center
        if col_idx in o_cols:
            cell.value = 'O'
            cell.font = font_regular

# Write Confirm rows
ws.merge_cells('A34:A52')
ws['A34'] = 'Confirm'
ws['A34'].fill = navy_fill
ws['A34'].font = font_header
ws['A34'].alignment = align_center_top

for r, b_val, c_val, o_cols in confirm_rows:
    if b_val:
        ws.cell(r, 2, b_val)
        if r in [34, 44, 47]:
            ws.cell(r, 2).font = font_bold
        else:
            ws.cell(r, 2).font = font_regular
    if c_val:
        ws.cell(r, 3, c_val)
        ws.cell(r, 3).font = font_regular
    
    for c in range(1, 4):
        ws.cell(r, c).border = thin_border
    
    for col_idx in range(4, 25):
        cell = ws.cell(r, col_idx)
        cell.border = thin_border
        cell.alignment = align_center
        if col_idx in o_cols:
            cell.value = 'O'
            cell.font = font_regular

# Section 3: Result (Rows 53-56)
ws.merge_cells('A53:A56')
ws['A53'] = 'Result'
ws['A53'].fill = navy_fill
ws['A53'].font = font_header
ws['A53'].alignment = align_center_top

ws['B53'] = 'Type (N : Normal, A : Abnormal, B : Boundary)'
ws['B53'].font = font_regular
ws['B54'] = 'Passed/Failed'
ws['B54'].font = font_regular
ws['B55'] = 'Executed Date'
ws['B55'].font = font_regular
ws['B56'] = 'Defect ID'
ws['B56'].font = font_regular

for r in range(53, 57):
    for c in range(1, 4):
        ws.cell(r, c).border = thin_border

# Results mapping: (Type, P/F, Date, DefectID)
test_results = [
    # UTCID01 to UTCID21
    ('N', 'P', '19/10/2026', None),       # UTCID01: Display all active
    ('A', 'F', '19/10/2026', 'SEARCH-D01'), # UTCID02: Display with inactive (Defect: displays inactive items)
    ('N', 'P', '19/10/2026', None),       # UTCID03: Exact full name
    ('N', 'P', '19/10/2026', None),       # UTCID04: Exact VN name
    ('N', 'P', '19/10/2026', None),       # UTCID05: Case-insensitive
    ('N', 'P', '19/10/2026', None),       # UTCID06: Partial suffix
    ('N', 'P', '19/10/2026', None),       # UTCID07: Partial prefix
    ('N', 'P', '19/10/2026', None),       # UTCID08: Partial multiple
    ('A', 'P', '19/10/2026', None),       # UTCID09: Non-existent name
    ('A', 'P', '19/10/2026', None),       # UTCID10: Random non-existent
    ('A', 'P', '19/10/2026', None),       # UTCID11: Deactivated product search
    ('B', 'P', '19/10/2026', None),       # UTCID12: Empty string
    ('B', 'F', '19/10/2026', 'SEARCH-D02'), # UTCID13: Leading/trailing space (Defect: untrimmed)
    ('B', 'P', '19/10/2026', None),       # UTCID14: Whitespace only
    ('B', 'P', '19/10/2026', None),       # UTCID15: Wildcard %
    ('A', 'P', '19/10/2026', None),       # UTCID16: SQL Injection
    ('A', 'P', '19/10/2026', None),       # UTCID17: XSS
    ('B', 'P', '19/10/2026', None),       # UTCID18: Max 150 chars
    ('B', 'P', '19/10/2026', None),       # UTCID19: Over max 151 chars
    ('A', 'P', '19/10/2026', None),       # UTCID20: Null parameter
    ('A', 'F', '19/10/2026', 'SEARCH-D03')  # UTCID21: DB down (Defect: uncaught NPE)
]

for i, (t_val, pf_val, date_val, defect_val) in enumerate(test_results):
    c = 4 + i
    # Row 53: Type
    cell53 = ws.cell(53, c, t_val)
    cell53.font = font_regular
    cell53.alignment = align_center
    cell53.border = thin_border
    
    # Row 54: P / F
    cell54 = ws.cell(54, c, pf_val)
    cell54.font = font_failed if pf_val == 'F' else font_bold
    cell54.alignment = align_center
    cell54.border = thin_border
    
    # Row 55: Date
    cell55 = ws.cell(55, c, date_val)
    cell55.font = font_regular
    cell55.alignment = align_rotate90
    cell55.border = thin_border
    
    # Row 56: Defect ID
    cell56 = ws.cell(56, c, defect_val)
    cell56.font = font_failed if defect_val else font_regular
    cell56.alignment = align_rotate90
    cell56.border = thin_border

# Defect List (Rows 58-61)
ws['A58'] = 'Defect list'
ws['A58'].font = font_bold

defects = [
    (59, 'SEARCH-D01', 'UTCID02: ProductJSP (listProduct) thực hiện "SELECT * FROM tblProducts" không có điều kiện "WHERE status = 1", dẫn đến sản phẩm đã ngừng hoạt động (status = 0) vẫn hiển thị trên danh sách với badge "DeActive", vi phạm yêu cầu "kết quả chỉ gồm sản phẩm đang hoạt động".'),
    (60, 'SEARCH-D02', 'UTCID13: ProductDAO.searchProductByName không trim khoảng trắng đầu/cuối của từ khóa tìm kiếm ("%" + productName + "%"), khiến câu truy vấn LIKE \'% Tao Fuji %\' không tìm thấy sản phẩm "Tao Fuji" và trả về danh sách rỗng.'),
    (61, 'SEARCH-D03', 'UTCID21: Khi mất kết nối cơ sở dữ liệu (connection = null), ProductDAO.searchProductByName và getAllProduct ném NullPointerException chưa được bắt (chỉ bắt SQLException), dẫn đến ứng dụng web báo lỗi HTTP 500.')
]

for r, d_id, d_desc in defects:
    ws.cell(r, 1, d_id).font = font_regular
    ws.cell(r, 2, d_desc).font = font_regular
    ws.cell(r, 2).alignment = align_left_wrap
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=24)

# Execution Notes (Rows 63-65)
ws['A63'] = 'Execution notes'
ws['A63'].font = font_bold

notes = [
    (64, 'Return / Exception: được kiểm thử tự động và thủ công bằng lớp kiểm thử JUnit ProductDAOTest và gọi trực tiếp ProductDAO trên cơ sở dữ liệu Assignment (UTCID01-21).'),
    (65, 'Screen message: được kiểm thử bằng cách gửi yêu cầu tìm kiếm và hiển thị trên ứng dụng web chạy trên máy chủ Apache Tomcat 10.1 (UTCID01-20). UTCID21 được xác nhận ở mức unit test.')
]

for r, n_desc in notes:
    ws.cell(r, 2, n_desc).font = font_regular
    ws.cell(r, 2).alignment = align_left_wrap
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=24)

print("Created searchProduct sheet.")

# Save workbook
wb.save(file_path)
print("Workbook saved successfully to:", file_path)
